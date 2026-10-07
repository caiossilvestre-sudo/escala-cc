"""Importa a planilha "Documentação IOT / NVRs" para o módulo Ferramentas.

Uso (dentro de backend/, com o venv ativo):

    # 1) SIMULAÇÃO — não grava nada, só gera o relatório de pendências
    python -m app.ferramentas.importar_planilha /caminho/planilha.xlsx

    # 2) Depois de revisar o relatório, grava de verdade
    python -m app.ferramentas.importar_planilha /caminho/planilha.xlsx --aplicar

Abas lidas:
    "NVR - Interno"            -> gravadores com origem "life"   (NVR instalado na Life)
    "NVR - Externos (Cliente)" -> gravadores com origem "cliente" (NVR instalado no cliente)
    "Life Guard"               -> câmeras LifeGuard (só linhas com ID da câmera)
As demais abas (resumos, contagens, cópias antigas, centrais de alarme)
ficam de fora nesta etapa.

Pode rodar de novo sem duplicar: gravador é localizado pelo nome, câmera
pelo gravador+canal (ou ID LifeGuard), credencial pelo usuário. O que já
existe no banco é mantido; só entra o que falta.

O relatório NUNCA contém senhas.
"""
import argparse
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv

load_dotenv()

import openpyxl  # noqa: E402

from app.ferramentas import normalizar as nz  # noqa: E402

ABAS_NVR = {"NVR - Interno": "life", "NVR - Externos (Cliente)": "cliente"}
ABA_LG = "Life Guard"


def _chave(h) -> str:
    s = unicodedata.normalize("NFD", str(h or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").upper()
    return re.sub(r"[^A-Z0-9]+", " ", s).strip()


# nome lógico -> possíveis cabeçalhos (já passados por _chave)
COLUNAS_NVR = {
    "gravador": ["GRAVADOR", "NOME NVR"],
    "porta_servico": ["PORTA DE SERVICO"],
    "url_acesso": ["IP DE ACESSO", "IP DE ACESSO NVR"],
    "canal": ["CANAL"],
    "status": ["STATUS CANAL"],
    "dias": ["DIAS DE ARMZ", "COLUNA1"],
    "nome": ["NOME CAM NVR"],
    "numero_cam": ["N CAM"],
    "descricao": ["DESCRICAO DO LOCAL", "DESCRICAO LOCAL CAM"],
    "id_cliente": ["ID CLIENTE IXC"],
    "cliente": ["NOME CLIENTE"],
    "contrato": ["CONTRATOS IXC", "CONTRATO IXC"],
    "cidade": ["CIDADE"],
    "pppoe": ["PPPOE DO CLIENTE"],
    "ip_pppoe": ["IP DO PPPOE"],
    "porta_publica": ["PORTA PUBLICA GRAVADOR"],
    "ip": ["IP CAMERA"],
    "porta": ["PORTA CAM", "PORTA CAMERA"],
    "modelo": ["MODELO CAMERA"],
    "compressao": ["TIPO DE COMPRESSAO"],
    "firmware": ["VERSAO FIRMWARE"],
    "mac": ["MAC CAMERA"],
    "obs": ["OBSERVACOES"],
}
COLUNAS_LG = {
    "ponto": ["NOME DO PONTO"],
    "contrato": ["CONTRATO CAM"],
    "endereco": ["ENDERECO"],
    "dias": ["DIAS DE ARMZ"],
    "contrato_vinculado": ["CONTRATO VINCULADO DADOS VPN"],
    "pppoe": ["PPPOE"],
    "ip_pppoe": ["IP PPPOE"],
    "lg_id": ["ID CAMERA"],
    "nome": ["NOME CAM"],
    "porta_lifeguard": ["PORTA LIFE GUARD"],
    "link": ["LINK PROVISIONAMENTO"],
    "ip": ["IP CAMERA"],
    "porta_monitoramento": ["PORTA MONITORAMENTO"],
    "modelo": ["MODELO CAMERA"],
    "compressao": ["TIPO DE COMPRESSAO"],
    "firmware": ["VERSAO FIRMWARE"],
    "mac": ["MAC CAMERA"],
}
# Campos do "ponto" LifeGuard que vêm em células mescladas: repete para as linhas de baixo
PREENCHER_LG = ["contrato", "endereco", "dias", "contrato_vinculado", "pppoe", "ip_pppoe"]


def _mapear(cabecalho, colunas):
    chaves = [_chave(h) for h in cabecalho]
    mapa = {}
    for nome, opcoes in colunas.items():
        for i, k in enumerate(chaves):
            if k in opcoes and nome not in mapa:
                mapa[nome] = i
    creds = []
    for i, k in enumerate(chaves):
        m = re.match(r"USUARIO DO GRAVADOR (\d+)$", k)
        if m:
            j = next((x for x, kk in enumerate(chaves) if kk == f"SENHA DO GRAVADOR {m.group(1)}"), None)
            creds.append((i, j))
    return mapa, creds


def _linhas(ws, max_col=70):
    for n, linha in enumerate(ws.iter_rows(max_col=max_col, values_only=True), start=1):
        yield n, linha


OCULTO = "(valor ocultado: igual a uma senha da planilha)"


class Relatorio:
    def __init__(self):
        self.pendencias = []
        self.contagem = Counter()
        self._vistos = set()
        self.senhas = set()  # preenchido antes da leitura; nada igual a uma senha sai no relatório

    def seguro(self, valor):
        t = nz.texto(valor)
        return OCULTO if t and t in self.senhas else t

    def add(self, aba, linha, referencia, campo, valor, problema, unico=False):
        """unico=True: registra só a 1ª ocorrência (mesmo gravador+valor+problema)."""
        if unico:
            chave = (aba, referencia, campo, str(valor), problema)
            if chave in self._vistos:
                return
            self._vistos.add(chave)
        valor = self.seguro(valor)
        self.pendencias.append((aba, linha, referencia, campo, "" if valor is None else str(valor)[:120], problema))
        self.contagem[problema] += 1

    def salvar(self, caminho: Path, resumo: dict):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Resumo"
        ws.append(["Item", "Quantidade"])
        for k, v in resumo.items():
            ws.append([k, v])
        ws.append([])
        ws.append(["Pendências por tipo", ""])
        for k, v in self.contagem.most_common():
            ws.append([k, v])
        ws.column_dimensions["A"].width = 60
        p = wb.create_sheet("Pendências")
        p.append(["Aba", "Linha", "Gravador / câmera", "Campo", "Valor na planilha", "Problema"])
        for row in self.pendencias:
            p.append(list(row))
        for col, w in zip("ABCDEF", (24, 8, 40, 16, 36, 60)):
            p.column_dimensions[col].width = w
        p.freeze_panes = "A2"
        p.auto_filter.ref = p.dimensions
        wb.save(caminho)


def _coletar_senhas(wb) -> set:
    senhas = set()
    for aba in ABAS_NVR:
        if aba not in wb.sheetnames:
            continue
        it = _linhas(wb[aba])
        _, cab = next(it)
        _, cols_cred = _mapear(cab, COLUNAS_NVR)
        idx = [j for _, j in cols_cred if j is not None]
        for _, linha in it:
            for j in idx:
                t = nz.texto(linha[j]) if j < len(linha) else None
                if t and len(t) >= 4:
                    senhas.add(t)
    return senhas


def ler_planilha(caminho: str, rel: Relatorio):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    rel.senhas = _coletar_senhas(wb)
    gravadores = {}                 # nome -> dict
    cameras_nvr = {}                # (nome_gravador, canal) -> dict
    credenciais = defaultdict(dict)  # nome_gravador -> {usuario: {"senha", "clientes"}}
    cameras_lg = {}                 # lg_id -> dict
    vazios = 0

    for aba, origem in ABAS_NVR.items():
        if aba not in wb.sheetnames:
            rel.add(aba, 0, "", "", "", "Aba não encontrada na planilha")
            continue
        it = _linhas(wb[aba])
        _, cab = next(it)
        mapa, cols_cred = _mapear(cab, COLUNAS_NVR)
        g = lambda linha, campo: linha[mapa[campo]] if campo in mapa and mapa[campo] < len(linha) else None  # noqa: E731
        for n, linha in it:
            nome_grav = nz.texto(g(linha, "gravador"))
            if not nome_grav:
                continue
            nome_grav = nome_grav.upper()
            ref = nome_grav
            grav = gravadores.setdefault(nome_grav, {
                "nome": nome_grav, "origem": origem, "url_acesso": None, "porta_servico": None,
                "dias": Counter(), "clientes": Counter(), "abas": set(),
            })
            grav["abas"].add(aba)
            if grav["origem"] != origem:
                rel.add(aba, n, ref, "gravador", nome_grav, "Gravador aparece nas abas Interno e Externos — ficou com a origem da primeira")
            grav["porta_servico"] = grav["porta_servico"] or nz.texto(g(linha, "porta_servico"))
            url = nz.texto(g(linha, "url_acesso"))
            if url and not grav["url_acesso"]:
                grav["url_acesso"] = url.replace(" ", "")

            # credenciais do gravador (repetidas em toda linha — junta e tira duplicadas)
            cliente_linha = nz.texto(g(linha, "cliente"))
            for iu, isenha in cols_cred:
                usuario = nz.texto(linha[iu]) if iu < len(linha) else None
                senha = nz.texto(linha[isenha]) if isenha is not None and isenha < len(linha) else None
                if not usuario and not senha:
                    continue
                if not usuario:
                    rel.add(aba, n, ref, "credencial", "", "Senha sem usuário — não importada", unico=True)
                    continue
                if nz.parece_endereco(usuario):
                    rel.add(aba, n, ref, "usuario", usuario, "Coluna de usuário com endereço — não importado como credencial", unico=True)
                    continue
                atual = credenciais[nome_grav].setdefault(usuario, {"senha": senha, "clientes": set()})
                if cliente_linha:
                    atual["clientes"].add(cliente_linha)
                if senha and atual["senha"] and senha != atual["senha"]:
                    rel.add(aba, n, ref, "usuario", usuario, "Mesmo usuário com senhas diferentes em linhas diferentes — mantida a primeira", unico=True)
                elif senha and not atual["senha"]:
                    atual["senha"] = senha

            canal = nz.inteiro(g(linha, "canal"))
            nome_cam = nz.texto(g(linha, "nome"))
            descricao = nz.texto(g(linha, "descricao"))
            ip_bruto, mac_bruto = g(linha, "ip"), g(linha, "mac")
            if not any(nz.texto(v) for v in (nome_cam, descricao, ip_bruto, mac_bruto)):
                vazios += 1  # canal sem câmera
                continue
            if canal is None:
                rel.add(aba, n, ref, "canal", g(linha, "canal"), "Canal vazio ou inválido — linha ignorada")
                continue
            ref = f"{nome_grav} · canal {canal}"
            if (nome_grav, canal) in cameras_nvr:
                rel.add(aba, n, ref, "canal", canal, "Canal repetido no mesmo gravador — mantida a primeira linha")
                continue

            obs = [rel.seguro(g(linha, "obs"))]
            ip_v, ip_ok = nz.ip(ip_bruto)
            if not ip_ok:
                rel.add(aba, n, ref, "ip", ip_bruto, "IP da câmera inválido — foi para Observações")
                obs.append(f"IP da câmera na planilha: {rel.seguro(ip_bruto)}")
                ip_v = None
            mac_v, mac_ok = nz.mac(mac_bruto)
            if not mac_ok:
                rel.add(aba, n, ref, "mac", mac_bruto, "MAC inválido — foi para Observações")
                obs.append(f"MAC na planilha: {rel.seguro(mac_bruto)}")
                mac_v = None
            ippp_v, ippp_ok = nz.ip(g(linha, "ip_pppoe"))
            if not ippp_ok:
                rel.add(aba, n, ref, "ip_pppoe", g(linha, "ip_pppoe"), "IP do PPPoE inválido — foi para Observações")
                obs.append(f"IP do PPPoE na planilha: {rel.seguro(g(linha, 'ip_pppoe'))}")
                ippp_v = None
            comp, comp_ok = nz.compressao(g(linha, "compressao"))
            if not comp_ok:
                rel.add(aba, n, ref, "compressao", "(valor ocultado)", "Valor estranho na coluna de compressão (pode ser senha anotada no lugar errado) — não importado")
            contratos = nz.contratos(g(linha, "contrato"))
            ids_cli = nz.contratos(g(linha, "id_cliente"))
            if len(contratos) > 1:
                obs.append(f"Contratos na planilha: {' / '.join(contratos)}")
                rel.add(aba, n, ref, "contrato", g(linha, "contrato"), "Mais de um contrato na célula — usado o primeiro, os outros foram para Observações")
            if len(ids_cli) > 1:
                obs.append(f"IDs de cliente na planilha: {' / '.join(ids_cli)}")
            # STATUS CANAL = ocupação do canal; câmera documentada com canal
            # "disponível" é sinal de câmera retirada ou planilha desatualizada
            if nz.status_canal(g(linha, "status")) == "disponivel":
                rel.add(aba, n, ref, "status", nz.texto(g(linha, "status")), "Canal marcado como DISPONÍVEL mas com dados de câmera — conferir se a câmera ainda existe")
                obs.append("Na planilha o canal estava marcado como disponível — conferir se a câmera ainda existe.")
            dias = nz.inteiro(g(linha, "dias"))
            if dias:
                grav["dias"][dias] += 1
            cliente = nz.texto(g(linha, "cliente"))
            if cliente:
                grav["clientes"][(cliente, ids_cli[0] if ids_cli else None, contratos[0] if contratos else None,
                                  nz.cidade(g(linha, "cidade")), nz.texto(g(linha, "pppoe")), ippp_v)] += 1

            cameras_nvr[(nome_grav, canal)] = {
                "tipo": "nvr", "canal": canal, "numero_cam": nz.texto(g(linha, "numero_cam")),
                "nome": nome_cam or descricao or f"CANAL {canal}",
                "descricao_local": descricao if descricao != nome_cam else None,
                "id_cliente_ixc": ids_cli[0] if ids_cli else None, "nome_cliente": cliente,
                "contrato_ixc": contratos[0] if contratos else None,
                "cidade": nz.cidade(g(linha, "cidade")), "pppoe": nz.texto(g(linha, "pppoe")),
                "ip_pppoe": ippp_v, "porta_publica": nz.texto(g(linha, "porta_publica")),
                "ip": ip_v, "porta": nz.texto(g(linha, "porta")), "mac": mac_v,
                "modelo": nz.texto(g(linha, "modelo")), "compressao": comp,
                "firmware": nz.texto(g(linha, "firmware")), "dias_gravacao": dias,
                "status": "desconhecido", "observacoes": "\n".join(o for o in obs if o) or None,
                "origem_importacao": f"{aba} (linha {n})",
            }
            if not nome_cam and not descricao:
                rel.add(aba, n, ref, "nome", "", "Câmera sem nome nem descrição — ficou como 'CANAL N'")

    # ---- LifeGuard
    if ABA_LG in wb.sheetnames:
        it = _linhas(wb[ABA_LG], max_col=25)
        _, cab = next(it)
        mapa, _ = _mapear(cab, COLUNAS_LG)
        g = lambda linha, campo: linha[mapa[campo]] if campo in mapa and mapa[campo] < len(linha) else None  # noqa: E731
        anterior, ponto_ant = {}, None
        sem_id = 0
        for n, linha in it:
            ponto = nz.texto(g(linha, "ponto"))
            if ponto != ponto_ant:
                anterior, ponto_ant = {}, ponto
            for campo in PREENCHER_LG:
                v = nz.texto(g(linha, campo))
                if v:
                    anterior[campo] = v
            lg_id = nz.texto(g(linha, "lg_id"))
            if not lg_id:
                if nz.texto(g(linha, "nome")):
                    sem_id += 1
                continue
            ref = f"LifeGuard {lg_id}"
            if lg_id in cameras_lg:
                rel.add(ABA_LG, n, ref, "lg_id", lg_id, "ID LifeGuard repetido — mantida a primeira linha")
                continue
            mac_v, mac_ok = nz.mac(g(linha, "mac"))
            if not mac_ok:
                rel.add(ABA_LG, n, ref, "mac", g(linha, "mac"), "MAC inválido — não importado")
                mac_v = None
            ip_v, ip_ok = nz.ip(g(linha, "ip"))
            if not ip_ok:
                rel.add(ABA_LG, n, ref, "ip", g(linha, "ip"), "IP da câmera inválido — não importado")
                ip_v = None
            ippp_v, ippp_ok = nz.ip(anterior.get("ip_pppoe"))
            if not ippp_ok:
                ippp_v = None
            comp, _ = nz.compressao(g(linha, "compressao"))
            contrato = (nz.contratos(anterior.get("contrato")) or [None])[0]
            vinculado = anterior.get("contrato_vinculado")
            obs = f"Contrato vinculado (dados/VPN): {vinculado}" if vinculado and vinculado != contrato else None
            cameras_lg[lg_id] = {
                "tipo": "lifeguard", "lg_id": lg_id, "nome": nz.texto(g(linha, "nome")) or f"LG {lg_id}",
                "descricao_local": anterior.get("endereco"), "nome_cliente": ponto,
                "contrato_ixc": contrato, "cidade": None, "pppoe": anterior.get("pppoe"), "ip_pppoe": ippp_v,
                "porta_lifeguard": nz.texto(g(linha, "porta_lifeguard")),
                "porta_monitoramento": nz.texto(g(linha, "porta_monitoramento")),
                "link_provisionamento": nz.texto(g(linha, "link")),
                "ip": ip_v, "mac": mac_v, "modelo": nz.texto(g(linha, "modelo")), "compressao": comp,
                "firmware": nz.texto(g(linha, "firmware")), "dias_gravacao": nz.inteiro(anterior.get("dias")),
                "status": "desconhecido", "observacoes": obs, "origem_importacao": f"{ABA_LG} (linha {n})",
            }
        if sem_id:
            rel.add(ABA_LG, 0, "", "lg_id", sem_id, f"{sem_id} linhas de câmera sem ID LifeGuard (pontos ainda não provisionados?) — não importadas")
    else:
        rel.add(ABA_LG, 0, "", "", "", "Aba não encontrada na planilha")

    # dados do gravador: cliente mais comum (para NVR externo é o dono do NVR) e dias de gravação mais comum
    for nome, grav in gravadores.items():
        dias = grav.pop("dias")
        grav["dias_gravacao"] = dias.most_common(1)[0][0] if dias else None
        cli = grav.pop("clientes")
        if grav["origem"] == "cliente" and cli:
            (cliente, id_cli, contrato, cidade, pppoe, ip_pppoe), _ = cli.most_common(1)[0]
            grav.update(nome_cliente=cliente, id_cliente_ixc=id_cli, contrato_ixc=contrato,
                        cidade=cidade, pppoe=pppoe, ip_pppoe=ip_pppoe)
        grav.pop("abas")
    return gravadores, cameras_nvr, credenciais, cameras_lg, vazios


def aplicar(gravadores, cameras_nvr, credenciais, cameras_lg):
    from app.db.session import Base, SessionLocal, engine
    from app.ferramentas.models import FtCamera, FtCredencial, FtGravador
    from app.ferramentas.seguranca import cifrar

    Base.metadata.create_all(bind=engine, tables=[FtGravador.__table__, FtCamera.__table__, FtCredencial.__table__])
    db = SessionLocal()
    novos = Counter()
    try:
        ids = {}
        for nome, dados in gravadores.items():
            g = db.query(FtGravador).filter(FtGravador.nome == nome).first()
            if not g:
                g = FtGravador(**dados)
                db.add(g)
                db.flush()
                novos["gravadores"] += 1
            ids[nome] = g.id
        for nome, creds in credenciais.items():
            existentes = {c.usuario for c in db.query(FtCredencial).filter(FtCredencial.gravador_id == ids[nome])}
            for ordem, (usuario, info) in enumerate(creds.items(), start=1):
                if usuario in existentes:
                    continue
                # usuário que aparece para um único cliente é daquele cliente;
                # em vários clientes (ou NVR externo, que é de um cliente só) é geral
                dono = next(iter(info["clientes"])) if len(info["clientes"]) == 1 and gravadores[nome]["origem"] == "life" else None
                db.add(FtCredencial(gravador_id=ids[nome], usuario=usuario, nome_cliente=dono,
                                    senha_cifrada=cifrar(info["senha"]), ordem=ordem))
                novos["credenciais"] += 1
        for (nome, canal), dados in cameras_nvr.items():
            if db.query(FtCamera).filter(FtCamera.gravador_id == ids[nome], FtCamera.canal == canal).first():
                continue
            db.add(FtCamera(gravador_id=ids[nome], **dados))
            novos["cameras_nvr"] += 1
        for lg_id, dados in cameras_lg.items():
            if db.query(FtCamera).filter(FtCamera.lg_id == lg_id).first():
                continue
            db.add(FtCamera(**dados))
            novos["cameras_lifeguard"] += 1
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
    return novos


def main(argv=None):
    ap = argparse.ArgumentParser(description="Importa a planilha de documentação de NVRs/LifeGuard.")
    ap.add_argument("planilha")
    ap.add_argument("--aplicar", action="store_true", help="grava no banco (sem isto, só simula)")
    ap.add_argument("--relatorio", default=None, help="caminho do relatório .xlsx")
    args = ap.parse_args(argv)

    rel = Relatorio()
    gravadores, cameras_nvr, credenciais, cameras_lg, vazios = ler_planilha(args.planilha, rel)
    origem = Counter(g["origem"] for g in gravadores.values())
    resumo = {
        "Gravadores NVR Life (instalados na Life)": origem["life"],
        "Gravadores NVR externos (instalados no cliente)": origem["cliente"],
        "Câmeras em NVR Life": sum(1 for (n, _) in cameras_nvr if gravadores[n]["origem"] == "life"),
        "Câmeras em NVR externo": sum(1 for (n, _) in cameras_nvr if gravadores[n]["origem"] == "cliente"),
        "Câmeras LifeGuard": len(cameras_lg),
        "Credenciais de gravador (usuário único por NVR)": sum(len(c) for c in credenciais.values()),
        "Canais vazios ignorados (sem nome, IP, MAC ou descrição)": vazios,
        "Linhas com pendência": len(rel.pendencias),
    }
    print("\n== Resultado da leitura ==")
    for k, v in resumo.items():
        print(f"  {k}: {v}")
    print("\n== Pendências por tipo ==")
    for k, v in rel.contagem.most_common():
        print(f"  {v:5d}  {k}")

    destino = Path(args.relatorio or f"relatorio_importacao_{datetime.now():%Y%m%d_%H%M}.xlsx")
    rel.salvar(destino, resumo)
    print(f"\nRelatório salvo em: {destino.resolve()}")

    if not args.aplicar:
        print("\nSIMULAÇÃO — nada foi gravado. Rode de novo com --aplicar para importar.")
        return 0
    novos = aplicar(gravadores, cameras_nvr, credenciais, cameras_lg)
    print("\n== Gravado no banco (só o que ainda não existia) ==")
    for k in ("gravadores", "credenciais", "cameras_nvr", "cameras_lifeguard"):
        print(f"  {k}: {novos[k]}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
